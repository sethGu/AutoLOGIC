from __future__ import annotations

import csv
import json
import math
import os
import re
import shutil
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
SRC = Path(os.environ.get("AUTOLOGIC_BEST40_SOURCE", ROOT / "outputs" / "fig4_best40"))
OUT = Path(os.environ.get("AUTOLOGIC_CLEAN_OUTPUT", ROOT / "outputs" / "fig4_clean_delivery"))
TABLES = OUT / "tables"
LOGS = OUT / "logs"
OLD_WORKBOOK = Path(os.environ.get("AUTOLOGIC_PAPER_WORKBOOK", ROOT / "data" / "paper" / "autoLOGIC_results_template.xlsx"))
NEW_WORKBOOK = OUT / "autoLOGIC_full_40_best_results_clean.xlsx"


TASK_SHEET = {
    "classification": "Classification",
    "regression": "Regression",
    "clustering": "Clustering",
}

SHEET_TASK = {v: k for k, v in TASK_SHEET.items()}

SECTIONS = {
    "Classification": [("AUC (%)", "auc"), ("ACC (%)", "acc")],
    "Regression": [("MAE", "mae"), ("RMSE", "rmse")],
    "Clustering": [("ARI (%)", "ari"), ("NMI (%)", "nmi")],
}

LOG_METRICS = {
    "classification": [("AUC", "auc"), ("ACC", "acc")],
    "regression": [("MAE", "mae"), ("RMSE", "rmse")],
    "clustering": [("ARI", "ari"), ("NMI", "nmi")],
}


def safe_float(value):
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    s = str(value).strip()
    if not s or s.lower() == "nan":
        return None
    try:
        return float(s)
    except Exception:
        return None


def fmt_value(value, digits=2):
    x = safe_float(value)
    if x is None:
        return ""
    return f"{x:.{digits}f}"


def fmt_avgstd(mean, std, digits=2):
    m = safe_float(mean)
    s = safe_float(std)
    if m is None or s is None:
        return ""
    return f"{m:.{digits}f}±{s:.{digits}f}"


def safe_name(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", text)


def write_csv(path: Path, rows: list[dict], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def update_main_workbook(summary: pd.DataFrame) -> None:
    shutil.copy2(OLD_WORKBOOK, NEW_WORKBOOK)
    wb = load_workbook(NEW_WORKBOOK)
    lookup = {(str(r.task), str(r.dataset)): r for r in summary.itertuples(index=False)}

    # Keep only the three original task sheets in the final delivery workbook.
    for sheet in list(wb.sheetnames):
        if sheet not in {"Classification", "Regression", "Clustering"}:
            del wb[sheet]

    for sheet, specs in SECTIONS.items():
        ws = wb[sheet]
        task = SHEET_TASK[sheet]
        for section_label, metric in specs:
            label_row = None
            for row in range(1, ws.max_row + 1):
                if ws.cell(row=row, column=1).value == section_label:
                    label_row = row
                    break
            if label_row is None:
                continue

            header_row = label_row + 1
            autologic_col = None
            for col in range(1, ws.max_column + 1):
                if ws.cell(row=header_row, column=col).value == "Auto-LOGIC":
                    autologic_col = col
                    break
            if autologic_col is None:
                continue

            row = header_row + 1
            while row <= ws.max_row and ws.cell(row=row, column=1).value:
                dataset = str(ws.cell(row=row, column=1).value).strip()
                record = lookup.get((task, dataset))
                if record is not None:
                    mean = getattr(record, f"{metric}_mean")
                    std = getattr(record, f"{metric}_std")
                    ws.cell(row=row, column=autologic_col).value = fmt_avgstd(mean, std)
                row += 1

    wb.save(NEW_WORKBOOK)


def write_clean_logs(seed_records: pd.DataFrame) -> list[dict]:
    manifest = []
    for (task, dataset), group in seed_records.groupby(["task", "dataset"], sort=True):
        group = group.sort_values("seed")
        metrics = LOG_METRICS[task]
        fields = ["dataset", "seed"] + [display for display, _ in metrics]
        rows = []
        for r in group.itertuples(index=False):
            row = {"dataset": dataset, "seed": int(getattr(r, "seed"))}
            for display, col in metrics:
                row[display] = fmt_value(getattr(r, col, None), 6)
            rows.append(row)

        file_name = safe_name(f"{task}_{dataset}_5seeds.log")
        path = LOGS / file_name
        write_csv(path, rows, fields)
        manifest.append({
            "task": task,
            "dataset": dataset,
            "log_file": str(path),
            "seed_count": len(rows),
        })
    return manifest


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    TABLES.mkdir(parents=True, exist_ok=True)
    LOGS.mkdir(parents=True, exist_ok=True)

    summary = pd.read_csv(SRC / "tables" / "per_dataset_best_setting_summary.csv")
    seed_records = pd.read_csv(SRC / "tables" / "per_seed_selected_records.csv")

    update_main_workbook(summary)
    log_manifest = write_clean_logs(seed_records)

    avgstd_fields = [
        "task", "dataset", "selection_mode", "setting", "seeds",
        "auc_mean", "auc_std", "acc_mean", "acc_std",
        "mae_mean", "mae_std", "rmse_mean", "rmse_std",
        "ari_mean", "ari_std", "nmi_mean", "nmi_std",
    ]
    write_csv(
        TABLES / "avgstd_values_used_for_main_excel.csv",
        summary.fillna("").to_dict("records"),
        avgstd_fields,
    )

    metric_fields = ["task", "dataset", "seed", "auc", "acc", "mae", "rmse", "ari", "nmi"]
    write_csv(
        TABLES / "per_seed_metrics_used_for_logs.csv",
        seed_records.fillna("").to_dict("records"),
        metric_fields,
    )
    write_csv(TABLES / "clean_logs_manifest.csv", log_manifest, ["task", "dataset", "log_file", "seed_count"])

    mixed = summary[summary["selection_mode"] != "strict_same_setting_5seeds"]
    manifest = {
        "output_dir": str(OUT),
        "main_workbook": str(NEW_WORKBOOK),
        "logs_dir": str(LOGS),
        "main_workbook_sheets": ["Classification", "Regression", "Clustering"],
        "dataset_count": int(summary.shape[0]),
        "seed_record_count": int(seed_records.shape[0]),
        "logs_count": len(log_manifest),
        "all_logs_have_5_seed_rows": all(x["seed_count"] == 5 for x in log_manifest),
        "autologic_column_updated_from": str(SRC / "tables" / "per_dataset_best_setting_summary.csv"),
        "strict_same_setting_dataset_count": int((summary["selection_mode"] == "strict_same_setting_5seeds").sum()),
        "mixed_available_seed_dataset_count": int(mixed.shape[0]),
        "mixed_available_seed_datasets": mixed[["task", "dataset", "setting", "seeds"]].to_dict("records"),
        "log_content_rule": "Each log contains only dataset, seed, and task performance metric columns.",
    }
    (OUT / "run_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
