from __future__ import annotations

import argparse
import math
import re
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import seaborn as sns
from openpyxl import load_workbook
from PIL import Image, ImageDraw


REPO_ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = REPO_ROOT / "data" / "paper" / "autoLOGIC_full_40_best_results_clean.xlsx"
OUT_DIR = REPO_ROOT / "outputs" / "fig4_subpanels"
CHECK_DIR = REPO_ROOT / "outputs" / "fig4_subpanels_checks"

DPI = 300
CANVAS_H_PX = 1080
PANEL_A_W_PX = 3500
PANEL_B_W_PX = 1600

AX_BOTTOM = 0.25
AX_HEIGHT = 0.62
AX_TOP = AX_BOTTOM + AX_HEIGHT
PANEL_A_AX = [0.095, AX_BOTTOM, 0.875, AX_HEIGHT]
PANEL_B_AX = [0.235, AX_BOTTOM, 0.725, AX_HEIGHT]

FONT = "Arial"
BLUE = "#2b83c6"
LIGHT_BLUE = "#93c9ff"
ORANGE = "#ff7f0e"
GREY = "#9a9a9a"
SPINE = "#5f5f5f"
GRID = "#e8e8e8"


SPECS = [
    {
        "sheet": "Classification",
        "section": "AUC (%)",
        "metric": "AUC",
        "task": "Classification",
        "baseline": "AutoGluon",
        "higher": True,
        "prefix": "01_classification_auc",
    },
    {
        "sheet": "Classification",
        "section": "ACC (%)",
        "metric": "ACC",
        "task": "Classification",
        "baseline": "AutoGluon",
        "higher": True,
        "prefix": "02_classification_acc",
    },
    {
        "sheet": "Regression",
        "section": "MAE",
        "metric": "MAE",
        "task": "Regression",
        "baseline": "AutoGluon",
        "higher": False,
        "prefix": "03_regression_mae",
    },
    {
        "sheet": "Regression",
        "section": "RMSE",
        "metric": "RMSE",
        "task": "Regression",
        "baseline": "AutoGluon",
        "higher": False,
        "prefix": "04_regression_rmse",
    },
    {
        "sheet": "Clustering",
        "section": "ARI (%)",
        "metric": "ARI",
        "task": "Clustering",
        "baseline": "GaussianMixture",
        "higher": True,
        "prefix": "05_clustering_ari",
    },
    {
        "sheet": "Clustering",
        "section": "NMI (%)",
        "metric": "NMI",
        "task": "Clustering",
        "baseline": "GaussianMixture",
        "higher": True,
        "prefix": "06_clustering_nmi",
    },
]


LABELS = {
    "RandomForest": "RF",
    "XGBoost": "XGB",
    "LightGBM": "LGBM",
    "TPOT": "TPOT",
    "H2O": "H2O",
    "AutoGluon": "AutoGluon",
    "DS-Agent": "DS-Agent",
    "RNE": "RNE",
    "Auto-LOGIC": "Auto-LOGIC",
    "Kmeans": "Kmeans",
    "GaussianMixture": "GMM",
    "AgglomerativeClustering": "Agglo",
}


def parse_mean(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("\u2212", "-")
    match = re.search(r"[-+]?\d+(?:\.\d+)?", text)
    return float(match.group(0)) if match else None


def find_section(ws, title: str):
    title_row = None
    for row in range(1, ws.max_row + 1):
        if ws.cell(row, 1).value == title:
            title_row = row
            break
    if title_row is None:
        raise ValueError(f"Cannot find section {title!r} in sheet {ws.title!r}")

    header_row = title_row + 1
    headers = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(header_row, col).value
        if val is not None:
            headers[str(val)] = col

    rows = []
    row = header_row + 1
    while row <= ws.max_row and ws.cell(row, 1).value is not None:
        dataset = str(ws.cell(row, 1).value)
        values = {}
        for method, col in headers.items():
            if method == "Dataset":
                continue
            val = parse_mean(ws.cell(row, col).value)
            if val is not None:
                values[method] = val
        rows.append({"dataset": dataset, "values": values})
        row += 1
    return headers, rows


def normalize_rows(rows, higher: bool):
    normalized = []
    for row in rows:
        vals = row["values"]
        finite = [v for v in vals.values() if v is not None and math.isfinite(v)]
        lo, hi = min(finite), max(finite)
        span = hi - lo
        norm_vals = {}
        for method, val in vals.items():
            if span == 0:
                norm = 1.0
            elif higher:
                norm = (val - lo) / span
            else:
                norm = (hi - val) / span
            norm_vals[method] = float(np.clip(norm, 0.0, 1.0))
        normalized.append({"dataset": row["dataset"], "values": norm_vals})
    return normalized


def wlt(rows, baseline: str, higher: bool):
    wins = losses = ties = 0
    for row in rows:
        auto = row["values"]["Auto-LOGIC"]
        base = row["values"][baseline]
        if abs(auto - base) < 1e-12:
            ties += 1
        elif (auto > base and higher) or (auto < base and not higher):
            wins += 1
        else:
            losses += 1
    return wins, losses, ties


def order_methods(norm_rows):
    buckets = {}
    for row in norm_rows:
        for method, value in row["values"].items():
            buckets.setdefault(method, []).append(value)
    return sorted(buckets, key=lambda m: (-float(np.mean(buckets[m])), LABELS.get(m, m)))


def setup_rc():
    plt.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": [FONT, "DejaVu Sans"],
            "axes.linewidth": 1.25,
            "xtick.major.width": 1.2,
            "ytick.major.width": 1.2,
            "xtick.major.size": 4.5,
            "ytick.major.size": 4.5,
            "savefig.dpi": DPI,
        }
    )


def save_panel_a(spec, norm_rows, out_path: Path):
    order = order_methods(norm_rows)
    xs, ys, methods = [], [], []
    for row in norm_rows:
        for method in order:
            if method in row["values"]:
                methods.append(LABELS.get(method, method))
                ys.append(row["values"][method])
                xs.append(method)

    fig = plt.figure(
        figsize=(PANEL_A_W_PX / DPI, CANVAS_H_PX / DPI),
        dpi=DPI,
        facecolor="white",
    )
    ax = fig.add_axes(PANEL_A_AX)

    sns.violinplot(
        x=[LABELS.get(m, m) for m in xs],
        y=ys,
        order=[LABELS.get(m, m) for m in order],
        ax=ax,
        inner=None,
        color=LIGHT_BLUE,
        saturation=0.55,
        linewidth=1.35,
        cut=0,
        width=0.86,
        density_norm="width",
        common_norm=False,
        bw_adjust=0.9,
    )
    for collection in ax.collections:
        collection.set_alpha(0.28)
        collection.set_edgecolor("#7dbfff")

    rng = np.random.default_rng(20260604)
    for i, method in enumerate(order):
        values = [
            row["values"][method]
            for row in norm_rows
            if method in row["values"]
        ]
        jitter = rng.uniform(-0.075, 0.075, size=len(values))
        ax.scatter(
            np.full(len(values), i) + jitter,
            values,
            s=18,
            color=LIGHT_BLUE,
            alpha=0.82,
            edgecolor="none",
            zorder=3,
        )
        mean = float(np.mean(values))
        std = float(np.std(values, ddof=1)) if len(values) > 1 else 0.0
        ax.errorbar(
            i,
            mean,
            yerr=std,
            fmt="o",
            color="#111111",
            ecolor="#111111",
            elinewidth=1.8,
            capsize=4.5,
            capthick=1.8,
            markersize=5.8,
            zorder=5,
        )

    ax.set_ylim(0, 1.0)
    ax.set_xlim(-0.52, len(order) - 0.48)
    ax.set_title(spec["task"], fontsize=19, weight="bold", pad=4)
    ax.set_ylabel(f"Normalized {spec['metric']}", fontsize=22, weight="bold", labelpad=12)
    ax.set_xlabel("")
    ax.set_yticks([0, 0.5, 1.0])
    ax.set_yticklabels(["0.0", "0.5", "1.0"], fontsize=16)
    ax.tick_params(axis="x", labelsize=15, pad=2)
    for label in ax.get_xticklabels():
        label.set_rotation(28)
        label.set_ha("right")
        label.set_rotation_mode("anchor")
    ax.grid(axis="y", color=GRID, linewidth=1.0)
    ax.grid(axis="x", visible=False)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_color(SPINE)
    ax.spines["bottom"].set_color(SPINE)

    fig.savefig(out_path, dpi=DPI, facecolor="white")
    plt.close(fig)


def save_panel_b(spec, raw_rows, norm_rows, out_path: Path):
    baseline = spec["baseline"]
    base_label = LABELS.get(baseline, baseline)
    wins, losses, ties = wlt(raw_rows, baseline, spec["higher"])

    xs = []
    ys = []
    colors = []
    for row in norm_rows:
        auto = row["values"]["Auto-LOGIC"]
        base = row["values"][baseline]
        xs.append(auto)
        ys.append(base)
        if abs(auto - base) < 1e-12:
            colors.append(GREY)
        elif auto > base:
            colors.append(BLUE)
        else:
            colors.append(ORANGE)

    fig = plt.figure(
        figsize=(PANEL_B_W_PX / DPI, CANVAS_H_PX / DPI),
        dpi=DPI,
        facecolor="white",
    )
    ax = fig.add_axes(PANEL_B_AX)

    ax.plot([0, 1], [0, 1], linestyle=":", color="black", linewidth=1.45, zorder=1)
    for x, y, c in zip(xs, ys, colors):
        ax.scatter(
            x,
            y,
            s=86,
            marker="o",
            facecolors="white",
            edgecolors=c,
            linewidths=2.9,
            zorder=3,
        )

    ax.annotate(
        "",
        xy=(0.02, 1.02),
        xytext=(0.31, 0.73),
        xycoords="axes fraction",
        textcoords="axes fraction",
        arrowprops=dict(arrowstyle="->", lw=2.25, color=ORANGE, mutation_scale=16),
        annotation_clip=False,
    )
    ax.annotate(
        "",
        xy=(0.98, 0.02),
        xytext=(0.70, 0.31),
        xycoords="axes fraction",
        textcoords="axes fraction",
        arrowprops=dict(arrowstyle="->", lw=2.25, color=BLUE, mutation_scale=16),
        annotation_clip=False,
    )

    ax.text(0.25, 0.60, f"{base_label}\nstronger", transform=ax.transAxes, fontsize=15)
    ax.text(0.58, 0.39, "Auto-LOGIC\nstronger", transform=ax.transAxes, fontsize=15)
    ax.text(
        0.055,
        0.075,
        f"W/L/T = {wins}/{losses}/{ties}",
        transform=ax.transAxes,
        fontsize=16,
        bbox=dict(facecolor="white", edgecolor="none", alpha=0.82, pad=1.0),
    )

    ax.set_xlim(-0.06, 1.06)
    ax.set_ylim(-0.06, 1.06)
    ticks = [0, 0.25, 0.5, 0.75, 1.0]
    ax.set_xticks(ticks)
    ax.set_yticks(ticks)
    ax.set_xticklabels([f"{t:.2f}" for t in ticks], fontsize=15)
    ax.set_yticklabels([f"{t:.2f}" for t in ticks], fontsize=15)
    ax.tick_params(axis="both", pad=2)
    ax.set_xlabel(f"Auto-LOGIC (norm {spec['metric']})", fontsize=18, labelpad=5)
    ax.set_ylabel(f"{base_label} (norm {spec['metric']})", fontsize=18, labelpad=7)

    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_color(SPINE)
    ax.spines["bottom"].set_color(SPINE)

    fig.savefig(out_path, dpi=DPI, facecolor="white")
    plt.close(fig)
    return wins, losses, ties


def pair_check(spec, panel_a: Path, panel_b: Path):
    a = Image.open(panel_a).convert("RGB")
    b = Image.open(panel_b).convert("RGB")
    gap = 48
    canvas = Image.new("RGB", (a.width + gap + b.width, CANVAS_H_PX), "white")
    canvas.paste(a, (0, 0))
    canvas.paste(b, (a.width + gap, 0))
    draw = ImageDraw.Draw(canvas)
    y0 = int(round(AX_BOTTOM * CANVAS_H_PX))
    y1 = int(round(AX_TOP * CANVAS_H_PX))
    for y in [y0, y1]:
        draw.line((0, y, canvas.width, y), fill=(220, 220, 220), width=2)
    canvas.save(CHECK_DIR / f"check_pair_{spec['prefix']}.png")


def image_bbox(path: Path):
    img = Image.open(path).convert("RGB")
    arr = np.asarray(img)
    mask = np.any(arr < 250, axis=2)
    ys, xs = np.where(mask)
    return {
        "width": img.width,
        "height": img.height,
        "bbox_left": int(xs.min()),
        "bbox_top": int(ys.min()),
        "bbox_right": int(xs.max()),
        "bbox_bottom": int(ys.max()),
    }


def main():
    global WORKBOOK, OUT_DIR, CHECK_DIR
    parser = argparse.ArgumentParser()
    parser.add_argument("panels", nargs="*", help="Optional panel prefixes to regenerate.")
    parser.add_argument("--workbook", type=Path, default=WORKBOOK)
    parser.add_argument("--out-dir", type=Path, default=OUT_DIR)
    parser.add_argument("--check-dir", type=Path, default=CHECK_DIR)
    args = parser.parse_args()
    WORKBOOK = args.workbook.expanduser().resolve()
    OUT_DIR = args.out_dir.expanduser().resolve()
    CHECK_DIR = args.check_dir.expanduser().resolve()
    setup_rc()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    CHECK_DIR.mkdir(parents=True, exist_ok=True)

    requested = set(args.panels)
    specs = [spec for spec in SPECS if not requested or spec["prefix"] in requested]
    if requested and len(specs) != len(requested):
        known = ", ".join(spec["prefix"] for spec in SPECS)
        raise ValueError(f"Unknown prefix in {sorted(requested)}. Known prefixes: {known}")

    if not requested:
        for old in OUT_DIR.glob("*.png"):
            old.unlink()
        for old in CHECK_DIR.glob("*.png"):
            old.unlink()
    else:
        for spec in specs:
            for old in (
                OUT_DIR / f"panel_a_{spec['prefix']}.png",
                OUT_DIR / f"panel_b_{spec['prefix']}.png",
                CHECK_DIR / f"check_pair_{spec['prefix']}.png",
            ):
                if old.exists():
                    old.unlink()

    wb = load_workbook(WORKBOOK, data_only=True)
    report = []

    for spec in specs:
        ws = wb[spec["sheet"]]
        _, rows = find_section(ws, spec["section"])
        norm_rows = normalize_rows(rows, spec["higher"])
        panel_a_path = OUT_DIR / f"panel_a_{spec['prefix']}.png"
        panel_b_path = OUT_DIR / f"panel_b_{spec['prefix']}.png"
        save_panel_a(spec, norm_rows, panel_a_path)
        wins, losses, ties = save_panel_b(spec, rows, norm_rows, panel_b_path)
        pair_check(spec, panel_a_path, panel_b_path)
        report.append(
            {
                "metric": spec["prefix"],
                "wlt": f"{wins}/{losses}/{ties}",
                "panel_a": image_bbox(panel_a_path),
                "panel_b": image_bbox(panel_b_path),
            }
        )

    for item in report:
        print(item["metric"], item["wlt"], item["panel_a"], item["panel_b"])
    print("OUT_DIR", OUT_DIR)
    print("CHECK_DIR", CHECK_DIR)


if __name__ == "__main__":
    main()
